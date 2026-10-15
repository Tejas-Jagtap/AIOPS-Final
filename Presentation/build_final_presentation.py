"""
Build the FINAL (End-Semester Defense) presentation for the AIOps project.

Design goals (fixes the problems of create_final_ppt.py):
  - Built from a BLANK presentation: no inherited corporate template,
    no "KONE Internal" watermark.
  - Strict layout system: every slide has a fixed content zone between the
    title rule and the footer band; images are LETTERBOXED into bounding
    boxes using their true pixel aspect ratio, so nothing ever overflows.
  - Left-aligned text (no justified-text rivers), inline **bold** emphasis.
  - COEP branding: navy/gold palette, logo top-right, slim footer band
    with department text and slide number that never collide.

Numbers come from evaluation/metrics/all_metrics.json (authoritative,
same as the journal paper).

Run:  python Presentation/build_final_presentation.py
Output: Presentation/AIOps_Final_Presentation.pptx  (25 slides, matches
        FINAL_PRESENTATION_SCRIPT.md slide-for-slide)
"""

import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------------- paths
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLOTS = os.path.join(BASE, 'evaluation', 'plots')
LOGO = os.path.join(BASE, 'Report', 'newcoepslogo.png')
ARCH = os.path.join(BASE, 'Report', 'architecutre.png')
OUT = os.path.join(BASE, 'Presentation', 'AIOps_Final_Presentation.pptx')

# ---------------------------------------------------------------- palette
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x54, 0x96)
GOLD = RGBColor(0xBF, 0x90, 0x00)
INK = RGBColor(0x26, 0x26, 0x26)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1E, 0x7B, 0x34)
RED = RGBColor(0xC0, 0x00, 0x00)
ROW_A = RGBColor(0xE8, 0xED, 0xF7)   # table banding
ROW_B = RGBColor(0xF6, 0xF8, 0xFC)
CALLOUT_BG = RGBColor(0xEA, 0xF0, 0xFA)
HILITE_BG = RGBColor(0xFF, 0xF2, 0xCC)  # proposed-row highlight

FONT = 'Calibri'

# ---------------------------------------------------------------- geometry
SW, SH = 13.333, 7.5
MARGIN = 0.55
TITLE_Y, TITLE_H = 0.20, 0.78
RULE_Y = 1.04
CONTENT_Y = 1.30          # content zone top
CONTENT_BOT = 6.92        # content zone bottom (nothing may pass this)
FOOTER_Y = 7.07           # footer band top

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]

# ---------------------------------------------------------------- helpers

def rich(p, text, size, color=INK, bold=False, italic=False,
         bold_color=None, font=FONT):
    """Add runs to paragraph p, honouring **bold** markers in text."""
    for i, part in enumerate(text.split('**')):
        if part == '':
            continue
        r = p.add_run()
        r.text = part
        f = r.font
        f.size = Pt(size)
        f.name = font
        f.italic = italic
        seg_bold = bold or (i % 2 == 1)
        f.bold = seg_bold
        f.color.rgb = (bold_color or color) if (i % 2 == 1 and bold_color) \
            else color
        f.bold = seg_bold


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tf


def bullets(slide, x, y, w, h, items, size=18, sub_size=16,
            space_after=8, color=INK, bold_color=NAVY):
    """items: list of (level, text) or plain str (level 0).
    level 0 -> '•', level 1 -> '–' indented, level 9 -> no glyph (spacer/plain)."""
    tf = textbox(slide, x, y, w, h)
    first = True
    for item in items:
        lvl, text = item if isinstance(item, tuple) else (0, item)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        if text == '':
            p.space_after = Pt(2)
            continue
        if lvl == 0:
            rich(p, '▪  ', size, BLUE, bold=True)
            rich(p, text, size, color, bold_color=bold_color)
        elif lvl == 1:
            p.level = 1
            rich(p, '–  ', sub_size, GRAY)
            rich(p, text, sub_size, color, bold_color=bold_color)
        else:  # plain paragraph, no glyph
            rich(p, text, size, color, bold_color=bold_color)
    return tf


def fit_image(slide, path, x, y, max_w, max_h, center=True):
    """Letterbox image into (x, y, max_w, max_h) preserving aspect ratio."""
    if not os.path.exists(path):
        print(f'  WARNING missing image: {path}')
        return None
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    if center:
        x += (max_w - w) / 2
        y += (max_h - h) / 2
    return slide.shapes.add_picture(path, Inches(x), Inches(y),
                                    Inches(w), Inches(h))


def rect(slide, x, y, w, h, fill, line=None, shadow_off=True, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    if shadow_off:
        shape.shadow.inherit = False
    return shape


def callout(slide, x, y, w, h, text, size=16, bg=CALLOUT_BG, fg=NAVY,
            border=BLUE, align=PP_ALIGN.CENTER):
    box = rect(slide, x, y, w, h, bg, line=border, rounded=True)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]
    p.alignment = align
    rich(p, text, size, fg, bold_color=fg)
    return box


def table(slide, x, y, w, col_widths, headers, data, header_size=13,
          data_size=12, row_h=0.38, header_h=0.42, cell_colors=None,
          row_fills=None, align_center_cols=()):
    """cell_colors: dict {(r,c): RGBColor} for text colour overrides.
    row_fills: dict {row_idx: RGBColor} fill override (data rows, 1-based)."""
    rows = len(data) + 1
    gfx = slide.shapes.add_table(rows, len(headers), Inches(x), Inches(y),
                                 Inches(w), Inches(header_h + row_h * len(data)))
    tbl = gfx.table
    tbl.first_row = False
    tbl.horz_banding = False
    for c, cw in enumerate(col_widths):
        tbl.columns[c].width = Inches(cw)
    tbl.rows[0].height = Inches(header_h)
    for r in range(1, rows):
        tbl.rows[r].height = Inches(row_h)
    for r in range(rows):
        for c in range(len(headers)):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
                p.alignment = PP_ALIGN.CENTER if c in align_center_cols \
                    or True else PP_ALIGN.LEFT
                rich(p, headers[c], header_size, WHITE, bold=True)
            else:
                fill = (row_fills or {}).get(r) or (ROW_A if r % 2 else ROW_B)
                cell.fill.solid()
                cell.fill.fore_color.rgb = fill
                val = data[r - 1][c]
                p.alignment = PP_ALIGN.CENTER if c in align_center_cols \
                    else PP_ALIGN.LEFT
                color = (cell_colors or {}).get((r, c), INK)
                rich(p, val, data_size, color, bold_color=NAVY)
    return tbl


def footer(slide, num):
    rect(slide, 0, FOOTER_Y, SW, SH - FOOTER_Y, NAVY)
    tf = textbox(slide, 0.4, FOOTER_Y, SW - 1.5, SH - FOOTER_Y,
                 anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    rich(p, 'Department of Computer Science and Engineering  |  '
            'COEP Technological University, Pune', 10.5, WHITE)
    tf2 = textbox(slide, SW - 1.0, FOOTER_Y, 0.7, SH - FOOTER_Y,
                  anchor=MSO_ANCHOR.MIDDLE)
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    rich(p2, str(num), 11, WHITE, bold=True)


def header(slide, title_text, num):
    """Standard body-slide chrome: title, gold rule, logo, footer."""
    tf = textbox(slide, MARGIN, TITLE_Y, 11.55, TITLE_H,
                 anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    rich(p, title_text, 29, NAVY, bold=True)
    rect(slide, MARGIN, RULE_Y, SW - 2 * MARGIN, 0.035, GOLD)
    fit_image(slide, LOGO, 12.42, 0.12, 0.70, 0.86)
    footer(slide, num)


def add_slide():
    return prs.slides.add_slide(BLANK)


n = 0

# =================================================================
# SLIDE 1 - Title
# =================================================================
n += 1
s = add_slide()
rect(s, 0, 0, SW, 1.32, NAVY)
tf = textbox(s, 0.5, 0.14, SW - 1.0, 1.1, anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'COEP Technological University, Pune', 22, WHITE, bold=True)
p = tf.add_paragraph()
p.alignment = PP_ALIGN.CENTER
rich(p, 'Department of Computer Science and Engineering', 14, WHITE)
rect(s, 0, 1.32, SW, 0.045, GOLD)

fit_image(s, LOGO, (SW - 1.05) / 2, 1.62, 1.05, 1.28)

tf = textbox(s, 0.9, 3.10, SW - 1.8, 1.35, anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'Intelligent AIOps Framework for Predictive CI/CD Pipeline '
        'Management and Autonomous Resource Optimization', 27, NAVY,
     bold=True)

tf = textbox(s, 0.9, 4.55, SW - 1.8, 0.4)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'M.Tech Dissertation — Final Defense', 15, GRAY, italic=True)

tf = textbox(s, 0.9, 5.10, SW - 1.8, 1.3)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'Tejas Annasaheb Jagtap', 20, INK, bold=True)
p = tf.add_paragraph()
p.alignment = PP_ALIGN.CENTER
rich(p, 'M.Tech (Information Security)  •  MIS: 712442021', 14, GRAY)
p = tf.add_paragraph()
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(10)
rich(p, 'Guide: **Dr. V. K. Pachghare**', 16, INK, bold_color=NAVY)
footer(s, n)
print(f'  Slide {n}: Title')

# =================================================================
# SLIDE 2 - Contents
# =================================================================
n += 1
s = add_slide()
header(s, 'Contents', n)
items = [
    'Introduction & Motivation',
    'Literature Survey & Research Gap',
    'Problem Statement & Objectives',
    'Proposed System Architecture',
    'Methodology: Data & the Three ML Models',
    'Evaluation Results (XGBoost, Autoencoder, LSTM)',
    'System Integration: FastAPI Orchestrator',
    'Predictive Autoscaler & Decision Algorithm',
    'End-to-End Validation & Demonstration',
    'Comparison with Related Work',
    'Conclusion, Future Scope & Publications',
]
half = 6
for col, chunk, x in ((0, items[:half], MARGIN + 0.35),
                      (1, items[half:], 7.0)):
    tf = textbox(s, x, 1.75, 5.8, 4.9)
    first = True
    for i, it in enumerate(chunk):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(16)
        idx = i + 1 + col * half
        rich(p, f'{idx:>2}.  ', 19, GOLD, bold=True)
        rich(p, it, 19, INK)
footer(s, n)
print(f'  Slide {n}: Contents')

# =================================================================
# SLIDE 3 - Introduction
# =================================================================
n += 1
s = add_slide()
header(s, 'Introduction', n)
bullets(s, MARGIN, CONTENT_Y + 0.1, SW - 2 * MARGIN, 4.1, [
    'Modern software delivery depends on **CI/CD pipelines (Jenkins)** and '
    'container orchestration with **Kubernetes**.',
    'These systems generate large volumes of build logs and infrastructure '
    'metrics every day.',
    'Traditional monitoring is **reactive** — it responds only after '
    'a failure or overload has already occurred:',
    (1, 'Build failures are discovered only after the build runs and wastes '
        'compute.'),
    (1, 'Log issues are found by manually reading thousands of log lines.'),
    (1, 'Kubernetes HPA adds pods only after CPU has already crossed a '
        'threshold.'),
], size=19, sub_size=17, space_after=12)
callout(s, 1.4, 5.85, SW - 2.8, 0.75,
        '**AIOps** (AI for IT Operations): apply Machine Learning to make '
        'these operations **predictive and self-optimizing**.', size=17)
print(f'  Slide {n}: Introduction')

# =================================================================
# SLIDE 4 - Motivation
# =================================================================
n += 1
s = add_slide()
header(s, 'Motivation', n)
bullets(s, MARGIN, CONTENT_Y + 0.05, SW - 2 * MARGIN, 4.3, [
    'Frequent **build failures** cause deployment delays and wasted compute.',
    '**Manual root-cause analysis** of logs is slow, tedious and error-prone.',
    '**Reactive autoscaling** introduces scaling lag — over-provisioning '
    '(cost) or degradation (SLA breach).',
    'Need: one self-learning framework that:',
    (1, 'Predicts build failures early from historical build patterns.'),
    (1, 'Detects log anomalies automatically, in real time.'),
    (1, 'Forecasts resource demand and scales proactively.'),
    (1, 'Closes the loop — predictions drive automatic Kubernetes '
        'actions.'),
], size=19, sub_size=17, space_after=11)
callout(s, 1.4, 5.95, SW - 2.8, 0.7,
        'Prior research treats these as **three separate problems** — '
        'no unified, closed-loop system exists.', size=17)
print(f'  Slide {n}: Motivation')

# =================================================================
# SLIDE 5 - Literature Survey (1)
# =================================================================
n += 1
s = add_slide()
header(s, 'Literature Survey', n)
table(s, MARGIN, 1.55, SW - 2 * MARGIN,
      [0.55, 3.35, 4.1, 4.23],
      ['Sr.', 'Paper', 'Work Done', 'Limitation'],
      [
          ['1', 'ML-Based Autoscaling for Cloud Resource Orchestration '
                '(J. Grid Computing, 2024)',
           'Compares ML models for predicting cluster load for autoscaling.',
           'Autoscaling only; **no CI/CD or log integration**.'],
          ['2', 'Time-series K8s Autoscaling using Prophet + LSTM '
                '(Frontiers in CS, 2025)',
           'Hybrid Prophet + LSTM forecast for proactive scaling.',
           'No CI/CD signals or log anomalies; **scaling in isolation**.'],
          ['3', 'Deep Learning for Multivariate Time-Series Anomaly '
                'Detection (AI Review, 2025)',
           'Survey of autoencoders / VAEs for anomaly detection.',
           'Survey only; **no concrete DevOps application**.'],
      ], data_size=13, row_h=1.15, align_center_cols=(0,))
print(f'  Slide {n}: Literature Survey (1)')

# =================================================================
# SLIDE 6 - Literature Survey (2)
# =================================================================
n += 1
s = add_slide()
header(s, 'Literature Survey (contd.)', n)
table(s, MARGIN, 1.45, SW - 2 * MARGIN,
      [0.55, 3.35, 4.1, 4.23],
      ['Sr.', 'Paper', 'Work Done', 'Limitation'],
      [
          ['4', 'Towards Optimization of Anomaly Detection in DevOps '
                '(IST, 2023)',
           'Evaluates anomaly-detection techniques for DevOps monitoring.',
           'Runtime anomalies only; **no CI/CD build prediction**.'],
          ['5', 'AI-Enhanced Platform Engineering for CI/CD (2025)',
           'Proposes AI integration for CI/CD optimization.',
           '**Conceptual** — no trained models, no validation.'],
          ['6', 'Predictive Failure Analysis in CI Pipelines using '
                'Ensemble Trees (IJACS, 2025)',
           'Ensemble tree models for CI build-failure prediction.',
           'Limited dataset; **no anomaly detection or scaling**.'],
          ['7', 'Agentic AIOps: From Reactive to Autonomous Operations '
                '(2025)',
           'Framework for closed-loop / autonomous AIOps.',
           'Conceptual closed loop; **no concrete models shipped**.'],
      ], data_size=13, row_h=1.02, align_center_cols=(0,))
print(f'  Slide {n}: Literature Survey (2)')

# =================================================================
# SLIDE 7 - Research Gap
# =================================================================
n += 1
s = add_slide()
header(s, 'Research Gap', n)
bullets(s, MARGIN, CONTENT_Y + 0.1, SW - 2 * MARGIN, 4.2, [
    'Existing studies address each problem in **isolation**:',
    (1, 'Build-failure prediction, **or**'),
    (1, 'Log anomaly detection, **or**'),
    (1, 'Predictive autoscaling.'),
    'No single framework **unifies all three signal types**.',
    'Reactive Kubernetes HPA still dominates — **predictive scaling '
    'is rare** in practice.',
    'Very few works build a true **closed loop**: prediction → decision '
    '→ automatic Kubernetes action → feedback.',
], size=19, sub_size=17, space_after=12)
callout(s, 1.4, 5.9, SW - 2.8, 0.75,
        '**Gap:** an integrated, closed-loop AIOps framework linking '
        '**Jenkins → ML → Kubernetes** in one loop.', size=17)
print(f'  Slide {n}: Research Gap')

# =================================================================
# SLIDE 8 - Problem Statement
# =================================================================
n += 1
s = add_slide()
header(s, 'Problem Statement', n)
box = rect(s, 1.15, 2.35, SW - 2.3, 2.5, CALLOUT_BG, line=BLUE, rounded=True)
tf = box.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tf.margin_left = tf.margin_right = Inches(0.45)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, '“To design and implement an intelligent, closed-loop AIOps '
        'framework that proactively **predicts CI/CD pipeline failures**, '
        '**detects anomalies** in Jenkins logs in real time, **forecasts** '
        'infrastructure workload, and autonomously optimizes Kubernetes '
        'resources through **predictive autoscaling**.”',
     22, INK, bold_color=NAVY)
print(f'  Slide {n}: Problem Statement')

# =================================================================
# SLIDE 9 - Objectives
# =================================================================
n += 1
s = add_slide()
header(s, 'Research Objectives', n)
objectives = [
    'Develop a **data-collection framework** for CI/CD build data and '
    'cluster metrics.',
    'Build **ML models** for failure prediction, log anomaly detection '
    'and CPU forecasting.',
    'Implement a **FastAPI-based AIOps orchestrator** serving all models '
    'via unified APIs.',
    'Design a **predictive autoscaling engine** for Kubernetes driven by '
    'ML outputs.',
    'Integrate the components into a **closed loop** and validate the '
    'system end-to-end.',
]
tf = textbox(s, MARGIN + 0.2, CONTENT_Y + 0.25, SW - 2 * MARGIN - 0.4, 4.0)
first = True
for i, ob in enumerate(objectives, 1):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(14)
    rich(p, f'{i}.  ', 20, GOLD, bold=True)
    rich(p, ob, 20, INK, bold_color=NAVY)
    rich(p, '   ✓', 20, GREEN, bold=True)
callout(s, 2.4, 6.0, SW - 4.8, 0.68,
        '**All five objectives completed** — verified by a 13/13 '
        'end-to-end test suite.', size=17, border=GREEN, fg=GREEN,
        bg=RGBColor(0xE9, 0xF5, 0xEC))
print(f'  Slide {n}: Objectives')

# =================================================================
# SLIDE 10 - Architecture
# =================================================================
n += 1
s = add_slide()
header(s, 'Proposed System Architecture', n)
fit_image(s, ARCH, 0.35, CONTENT_Y, 4.15, CONTENT_BOT - CONTENT_Y)
bullets(s, 4.95, CONTENT_Y + 0.05, SW - 4.95 - MARGIN, 5.6, [
    'Closed-loop pipeline with **four layers**:',
    '**1. Data Collectors**',
    (1, 'Jenkins REST API → build metadata + console logs.'),
    (1, 'Prometheus HTTP API → CPU / memory metrics.'),
    '**2. ML Layer** — three models:',
    (1, '**XGBoost** → build-failure prediction.'),
    (1, '**TF-IDF + Autoencoder** → log anomaly detection.'),
    (1, '**LSTM** → CPU workload forecasting.'),
    '**3. AIOps Orchestrator (FastAPI)** — loads all models, exposes '
    'REST endpoints.',
    '**4. Predictive Autoscaler** — decision engine → Kubernetes '
    'Python client.',
    'Scaling actions change cluster state → new metrics **feed back** '
    'into the next cycle.',
], size=17, sub_size=15.5, space_after=7)
print(f'  Slide {n}: Architecture')

# =================================================================
# SLIDE 11 - Methodology: data
# =================================================================
n += 1
s = add_slide()
header(s, 'Methodology: Data Pipeline & Collection', n)
tf = textbox(s, MARGIN, CONTENT_Y + 0.05, 6.0, 0.45)
rich(tf.paragraphs[0], 'Synthetic training dataset', 18, NAVY, bold=True)
bullets(s, MARGIN, CONTENT_Y + 0.55, 6.0, 3.3, [
    '**600 Jenkins builds** — 68% SUCCESS, 23% FAILURE, 9% UNSTABLE.',
    '**600 console logs** with class-distinctive vocabulary.',
    '**10,080 CPU points** — daily / weekly patterns + noise + spikes.',
    'Reproducible, with **labelled ground truth** for evaluation.',
], size=16.5, space_after=10)
tf = textbox(s, 7.0, CONTENT_Y + 0.05, 5.8, 0.45)
rich(tf.paragraphs[0], 'Live data-collection infrastructure', 18, NAVY,
     bold=True)
bullets(s, 7.0, CONTENT_Y + 0.55, 5.8, 3.3, [
    '**Jenkins REST collector** — build metadata + console logs.',
    '**Prometheus collector** — CPU / memory queries.',
    '**Real-time streaming collector** with configurable interval.',
    'Same pipeline works against **real Jenkins / Prometheus** instances.',
], size=16.5, space_after=10)
rect(s, 6.62, CONTENT_Y + 0.15, 0.012, 3.5, RGBColor(0xC8, 0xD2, 0xE4))
callout(s, 1.0, 5.75, SW - 2.0, 0.85,
        '**Preprocessing:**  8 engineered build features  •  '
        'TF-IDF (500 dims, n-gram 1–2)  •  MinMax scaling + '
        '60-step windows for CPU.', size=16)
print(f'  Slide {n}: Methodology data')

# =================================================================
# SLIDES 12-14 - the three models (shared layout)
# =================================================================

def model_slide(num, title, desc_items, metric_rows, img1, img2,
                note=None):
    s = add_slide()
    header(s, title, num)
    bullets(s, MARGIN, CONTENT_Y + 0.05, 6.7, 2.8, desc_items,
            size=15.5, sub_size=14.5, space_after=5)
    table(s, 7.6, CONTENT_Y + 0.05, 5.15, [2.6, 2.55],
          ['Metric', 'Value'], metric_rows,
          header_size=12.5, data_size=12, row_h=0.34, header_h=0.36,
          align_center_cols=(1,))
    img_y, img_h = 4.28, 2.55
    iw1 = Image.open(img1).size
    iw2 = Image.open(img2).size
    w1 = img_h * iw1[0] / iw1[1]
    w2 = img_h * iw2[0] / iw2[1]
    gap = 0.45
    start = (SW - w1 - w2 - gap) / 2
    fit_image(s, img1, start, img_y, w1, img_h)
    fit_image(s, img2, start + w1 + gap, img_y, w2, img_h)
    if note:
        tf = textbox(s, MARGIN, 3.9, SW - 2 * MARGIN, 0.35)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        rich(p, note, 13.5, GRAY, italic=True)
    print(f'  Slide {num}: {title}')
    return s

n += 1
model_slide(
    n, 'Model 1: XGBoost — Failure Prediction',
    [
        'Gradient-boosted ensemble of **100 trees**.',
        'Input: **8 build features** — duration, warnings, errors, '
        'failures, log length, tests failed / passed / total.',
        'Output: **failure probability** (0–1).',
        'Class imbalance handled via **scale_pos_weight**.',
        '80/20 stratified split + **5-fold cross-validation**.',
        'CV F1 = **0.975 ± 0.009** → stable, no overfitting.',
        '**Recall = 100%** — the model never misses a real failure.',
    ],
    [['Accuracy', '**99.2%**'], ['Precision', '96.4%'],
     ['Recall', '**100%**'], ['F1-Score', '**0.982**'],
     ['ROC-AUC', '1.000']],
    os.path.join(PLOTS, 'xgboost_confusion_matrix.png'),
    os.path.join(PLOTS, 'xgboost_feature_importance.png'),
    note='Feature importance confirms real signal: test results and error '
         'counts dominate — consistent with empirical CI studies.')

n += 1
model_slide(
    n, 'Model 2: TF-IDF + Autoencoder — Log Anomaly Detection',
    [
        'Two-stage, fully **unsupervised** pipeline (no labels needed).',
        '**Stage 1:** TF-IDF vectorization — 500 features, '
        'n-gram 1–2, sublinear TF.',
        '**Stage 2:** Autoencoder 500 → 128 → 64 → **32** '
        '→ 64 → 128 → 500.',
        'Trained on **normal logs only** — learns what “normal” '
        'looks like.',
        'Anomaly = **high reconstruction error** (MSE above threshold).',
        'Threshold = 95th percentile of normal reconstruction errors.',
    ],
    [['Accuracy', '**96.7%**'], ['Precision', '87.1%'],
     ['Recall', '**100%**'], ['F1-Score', '**0.931**'],
     ['ROC-AUC', '1.000']],
    os.path.join(PLOTS, 'anomaly_confusion_matrix.png'),
    os.path.join(PLOTS, 'anomaly_error_distribution.png'),
    note='Error distribution shows clean separation between normal and '
         'anomalous logs.')

n += 1
model_slide(
    n, 'Model 3: LSTM — CPU Workload Forecasting',
    [
        'Input: last **60 minutes** of CPU usage (sliding window).',
        'Output: CPU usage **10 minutes ahead** — the lead time the '
        'autoscaler needs.',
        'Architecture: **LSTM(128) → LSTM(64) → Dense(32, ReLU) '
        '→ Dense(1)** with dropout.',
        'MinMax normalization to [0, 1]; ~**10,000 sequences**, '
        'chronological split.',
        'Early stopping + learning-rate reduction on plateau.',
        'MAPE **9.68%** — under the 10% rule-of-thumb for good '
        'forecasts.',
        'Model tracks the **trend** and smooths unpredictable spikes.',
    ],
    [['MSE', '0.0049'], ['MAE', '0.0433'], ['RMSE', '**0.0698**'],
     ['MAPE', '**9.68%**'], ['R²', '0.455']],
    os.path.join(PLOTS, 'lstm_actual_vs_predicted.png'),
    os.path.join(PLOTS, 'lstm_scatter.png'))

# =================================================================
# SLIDE 15 - Combined performance
# =================================================================
n += 1
s = add_slide()
header(s, 'Combined Model Performance Summary', n)
table(s, 0.85, 1.5, SW - 1.7,
      [3.0, 3.4, 2.7, 2.53],
      ['Module', 'Algorithm', 'Key Metric', 'ROC-AUC / R²'],
      [
          ['Failure Prediction', 'XGBoost', '**F1 = 0.982**', 'AUC = 1.000'],
          ['Anomaly Detection', 'TF-IDF + Autoencoder', '**F1 = 0.931**',
           'AUC = 1.000'],
          ['CPU Forecasting', 'LSTM', '**RMSE = 0.070**', 'R² = 0.455'],
      ], data_size=13.5, row_h=0.48, align_center_cols=(2, 3))
fit_image(s, os.path.join(PLOTS, 'model_comparison_summary.png'),
          1.6, 3.55, SW - 3.2, 2.75)
tf = textbox(s, MARGIN, 6.42, SW - 2 * MARGIN, 0.4)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'Both classifiers reach **100% recall** — a deliberate design '
        'choice: a missed failure costs far more than a false alarm.',
     15, INK, bold_color=NAVY)
print(f'  Slide {n}: Combined summary')

# =================================================================
# SLIDE 16 - ROC curves
# =================================================================
n += 1
s = add_slide()
header(s, 'ROC Curves — Classification Models', n)
img_h = 4.35
w = img_h * 1.2
gap = 0.9
start = (SW - 2 * w - gap) / 2
fit_image(s, os.path.join(PLOTS, 'xgboost_roc_curve.png'),
          start, 1.45, w, img_h)
fit_image(s, os.path.join(PLOTS, 'anomaly_roc_curve.png'),
          start + w + gap, 1.45, w, img_h)
for x, label in ((start, 'XGBoost: AUC = 1.000'),
                 (start + w + gap, 'Autoencoder: AUC = 1.000')):
    tf = textbox(s, x, 5.9, w, 0.4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    rich(p, f'**{label}**', 17, NAVY)
tf = textbox(s, MARGIN, 6.45, SW - 2 * MARGIN, 0.4)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'AUC = 1.000 reflects the clean class separation of the synthetic '
        'dataset — acknowledged in the limitations and future work.',
     13.5, GRAY, italic=True)
print(f'  Slide {n}: ROC curves')

# =================================================================
# SLIDE 17 - Orchestrator
# =================================================================
n += 1
s = add_slide()
header(s, 'System Integration: FastAPI Orchestrator', n)
bullets(s, MARGIN, CONTENT_Y + 0.15, 6.4, 4.8, [
    'A **single FastAPI service** loads all three models into memory at '
    'startup and serves them via REST.',
    '**Async endpoints** with typed Pydantic request / response schemas.',
    'Models cached in memory → **millisecond-scale inference**.',
    'Live **Swagger UI** at /docs for interactive testing.',
    'All external endpoints come from **environment variables** — '
    'nothing hard-coded.',
    'Health checks + graceful startup / shutdown via lifespan events.',
], size=17.5, sub_size=16, space_after=12)
table(s, 7.35, CONTENT_Y + 0.35, 5.4, [2.9, 2.5],
      ['REST Endpoint', 'Backed by'],
      [
          ['POST  /predict_failure', 'XGBoost'],
          ['POST  /detect_anomaly', 'Autoencoder'],
          ['POST  /forecast_cpu', 'LSTM'],
          ['POST  /autoscale', 'Decision Engine'],
      ], data_size=13, row_h=0.52, header_h=0.44)
callout(s, 7.35, 5.35, 5.4, 1.0,
        'One unified API layer connects **models → decisions → '
        'Kubernetes actions**.', size=15)
print(f'  Slide {n}: Orchestrator')

# =================================================================
# SLIDE 18 - Autoscaler
# =================================================================
n += 1
s = add_slide()
header(s, 'Predictive Autoscaler — Decision Algorithm', n)
tf = textbox(s, MARGIN, CONTENT_Y + 0.05, 6.1, 0.4)
rich(tf.paragraphs[0], 'Priority-based policy (each scaling cycle)', 17,
     NAVY, bold=True)
box = rect(s, MARGIN, CONTENT_Y + 0.55, 6.1, 3.15, RGBColor(0xF4, 0xF6, 0xFB),
           line=BLUE, rounded=True)
tf = box.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.25)
tf.margin_top = Inches(0.15)
steps = [
    ('1.', 'Anomaly detected → **raise ALERT** (non-blocking).'),
    ('2.', 'Failure probability > 0.6 → **HOLD** — don’t '
           'scale a likely-failing build.'),
    ('3.', 'Predicted CPU ≥ 0.7 → **SCALE UP (+2 replicas)**.'),
    ('4.', 'Predicted CPU ≤ 0.3 → **SCALE DOWN (−1 replica)**.'),
    ('5.', 'Otherwise → **MAINTAIN**.'),
]
first = True
for num_, txt in steps:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(9)
    rich(p, f'{num_}  ', 16, GOLD, bold=True)
    rich(p, txt, 16, INK, bold_color=NAVY)
bullets(s, MARGIN, 5.15, 6.1, 1.6, [
    '**Guardrails:** 1–10 replicas, **300 s cooldown** against '
    'thrashing.',
    'Actions applied via the **Kubernetes Python client** → closes '
    'the loop.',
], size=15.5, space_after=8)
tf = textbox(s, 7.1, CONTENT_Y + 0.05, 5.65, 0.4)
rich(tf.paragraphs[0], 'Why predictive? (vs. Kubernetes HPA)', 17, NAVY,
     bold=True)
table(s, 7.1, CONTENT_Y + 0.55, 5.65, [1.55, 2.0, 2.1],
      ['Aspect', 'HPA (reactive)', 'Proposed'],
      [
          ['Trigger', 'Current CPU', '**Forecast (t + 10 min)**'],
          ['Timing', 'After the spike', '**Before the spike**'],
          ['Signals', 'CPU only', '**CPU + failures + log anomalies**'],
          ['CI/CD aware', 'No', '**Yes (HOLD rule)**'],
      ], data_size=12.5, row_h=0.62, header_h=0.42)
callout(s, 7.1, 5.6, 5.65, 0.95,
        'The 10-minute forecast gives Kubernetes **lead time** to add pods '
        '**before** the load arrives.', size=14.5)
print(f'  Slide {n}: Autoscaler')

# =================================================================
# SLIDE 19 - E2E validation
# =================================================================
n += 1
s = add_slide()
header(s, 'End-to-End Validation & Demonstration', n)
bullets(s, MARGIN, CONTENT_Y + 0.1, 7.3, 4.3, [
    'Automated end-to-end suite: **13 / 13 tests pass**.',
    '**Data layer** — dataset presence, class distribution, value '
    'ranges.',
    '**Model layer** — all six model artefacts load correctly.',
    '**API layer** — every endpoint returns correct results:',
    (1, 'SUCCESS build → failure probability **0.001**.'),
    (1, 'FAILURE build → failure probability **0.97**.'),
    (1, 'Normal log → pass;  failure log → **flagged anomaly**.'),
    (1, 'CPU forecast within valid range [0, 1].'),
    '**Full pipeline** — failure + anomaly + forecast fused into one '
    'scaling decision.',
], size=16.5, sub_size=15, space_after=8)
table(s, 8.35, CONTENT_Y + 0.35, 4.4, [2.45, 1.95],
      ['Test Group', 'Result'],
      [
          ['Data layer', '3 / 3  PASS'],
          ['Model layer', '2 / 2  PASS'],
          ['API layer', '7 / 7  PASS'],
          ['Full pipeline', '1 / 1  PASS'],
          ['**Total**', '**13 / 13**'],
      ], data_size=13, row_h=0.5, header_h=0.44, align_center_cols=(1,),
      cell_colors={(1, 1): GREEN, (2, 1): GREEN, (3, 1): GREEN,
                   (4, 1): GREEN, (5, 1): GREEN})
callout(s, 1.0, 6.05, SW - 2.0, 0.72,
        'Runs on real infrastructure: **Next.js demo app**, six-stage '
        '**Jenkinsfile**, **Docker** images and **Kubernetes manifests**.',
        size=15.5)
print(f'  Slide {n}: E2E validation')

# =================================================================
# SLIDE 20 - Comparison with related work
# =================================================================
n += 1
s = add_slide()
header(s, 'Comparison with Related Work', n)
CHK, CRO = '✓', '✗'
data = [
    ['DeepLog / log-based detection (Du et al.)', CRO, CHK, CRO, CRO],
    ['Guruge & Priyadarshana (2025)', CRO, CRO, CHK, CRO],
    ['Da Silva et al. (2025)', CRO, CRO, CHK, CRO],
    ['Zota et al. (2025, agentic AIOps)', CRO, CRO, CRO, CHK],
    ['**Proposed Framework**', CHK, CHK, CHK, CHK],
]
colors = {}
for r in range(1, 6):
    for c in range(1, 5):
        colors[(r, c)] = GREEN if data[r - 1][c] == CHK else RED
table(s, 0.7, 1.6, SW - 1.4,
      [4.35, 1.9, 1.9, 2.0, 1.78],
      ['Study', 'Failure Pred.', 'Log Anomaly', 'Predictive Scaling',
       'Closed Loop'],
      data, data_size=13.5, row_h=0.62, header_h=0.5,
      align_center_cols=(1, 2, 3, 4),
      cell_colors=colors, row_fills={5: HILITE_BG})
callout(s, 1.5, 5.65, SW - 3.0, 0.8,
        'The proposed framework is the **only one integrating all four '
        'capabilities** in a single closed-loop architecture.', size=16.5)
print(f'  Slide {n}: Comparison')

# =================================================================
# SLIDE 21 - Conclusion
# =================================================================
n += 1
s = add_slide()
header(s, 'Conclusion', n)
bullets(s, MARGIN, CONTENT_Y + 0.15, SW - 2 * MARGIN, 5.3, [
    'Built a **unified, closed-loop AIOps framework**: three ML models + '
    'FastAPI orchestrator + predictive Kubernetes autoscaler.',
    '**XGBoost** predicts build failures — **F1 = 0.982, 100% recall** '
    '(no missed failures).',
    '**TF-IDF + Autoencoder** detects log anomalies unsupervised — '
    '**F1 = 0.931, 100% recall**.',
    '**LSTM** forecasts CPU — **RMSE = 0.070, MAPE = 9.68%** — '
    'enabling proactive scaling.',
    'Predictive autoscaling acts on forecasts, overcoming the lag of '
    'reactive Kubernetes HPA.',
    'Validated **end-to-end (13/13 tests)** — a concrete implementation '
    'of the closed-loop AIOps that recent literature describes only '
    'conceptually.',
], size=18.5, space_after=14)
print(f'  Slide {n}: Conclusion')

# =================================================================
# SLIDE 22 - Future scope
# =================================================================
n += 1
s = add_slide()
header(s, 'Future Scope', n)
bullets(s, MARGIN, CONTENT_Y + 0.25, SW - 2 * MARGIN, 5.2, [
    'Validate on **real-world CI/CD datasets** (TravisTorrent, BuildFast) '
    'and live Jenkins / Prometheus data.',
    'Add **automatic retraining** to handle concept drift as pipelines '
    'evolve.',
    'Replace TF-IDF with **transformer-based log embeddings** (e.g. '
    'LogBERT) for richer context.',
    '**Benchmark against Kubernetes HPA** under production-scale load.',
    'Build a **monitoring dashboard** for predictions, alerts and scaling '
    'actions.',
    'Extend to **multi-service dependencies** and memory / network '
    'forecasting.',
], size=18.5, space_after=15)
print(f'  Slide {n}: Future scope')

# =================================================================
# SLIDE 23 - Publications
# =================================================================
n += 1
s = add_slide()
header(s, 'Publications', n)
callout(s, 1.0, 1.75, SW - 2.0, 1.15,
        '“Intelligent AIOps Framework for Predictive CI/CD Pipeline '
        'Management and Autonomous Resource Optimization”',
        size=19)
bullets(s, MARGIN + 0.4, 3.3, SW - 2 * MARGIN - 0.8, 3.2, [
    '**Authors:** Tejas A. Jagtap, Charulata Patil, V. K. Pachghare — '
    'Dept. of Computer Science & Engineering, COEP Technological '
    'University, Pune.',
    '**Journal manuscript** prepared in Springer Nature format; '
    '**conference version** also prepared.',
    '**Contribution:** first unified closed-loop AIOps framework combining '
    'failure prediction, log anomaly detection, CPU forecasting and '
    'predictive autoscaling.',
], size=18, space_after=16)
print(f'  Slide {n}: Publications')

# =================================================================
# SLIDE 24 - References
# =================================================================
n += 1
s = add_slide()
header(s, 'References', n)
refs = [
    'I. Pintye et al., “Enhancing ML-Based Autoscaling for Cloud '
    'Resource Orchestration,” Journal of Grid Computing, 2024.',
    'P. B. Guruge and Y. H. P. P. Priyadarshana, “Time-Series '
    'Forecasting-Based Kubernetes Autoscaling Using Prophet and '
    'LSTM,” Frontiers in Computer Science, 2025.',
    'M. A. Yahya et al., “Deep Learning for Multivariate Time-Series '
    'Anomaly Detection: A Survey,” Artificial Intelligence Review, '
    '2025.',
    'A. Hrusto, “Towards Optimization of Anomaly Detection in '
    'DevOps,” Information and Software Technology, 2023.',
    'P. Mathew and A. Baswareddy, “Predictive Failure Analysis in CI '
    'Pipelines Using Ensemble Trees,” IJACS, 2025.',
    'M. Du et al., “DeepLog: Anomaly Detection and Diagnosis from '
    'System Logs,” ACM CCS, 2017.',
    'Y. Zota et al., “A Practical Framework for Agentic AIOps,” '
    '2025.',
    'Kubernetes Documentation, “Horizontal Pod Autoscaler,” 2024.',
]
tf = textbox(s, MARGIN, CONTENT_Y + 0.15, SW - 2 * MARGIN, 5.3)
first = True
for i, ref in enumerate(refs, 1):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(11)
    rich(p, f'[{i}]  ', 15, NAVY, bold=True)
    rich(p, ref, 15, INK)
print(f'  Slide {n}: References')

# =================================================================
# SLIDE 25 - Thank you
# =================================================================
n += 1
s = add_slide()
rect(s, 0, 2.55, SW, 0.05, GOLD)
tf = textbox(s, 1.0, 2.85, SW - 2.0, 1.3, anchor=MSO_ANCHOR.MIDDLE)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'Thank You', 60, NAVY, bold=True)
rect(s, 0, 4.35, SW, 0.05, GOLD)
tf = textbox(s, 1.0, 4.7, SW - 2.0, 0.5)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'Questions & Discussion', 22, GRAY)
tf = textbox(s, 1.0, 5.45, SW - 2.0, 0.4)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
rich(p, 'Tejas Annasaheb Jagtap  •  M.Tech (Information Security)  '
        '•  COEP Technological University, Pune', 14, GRAY)
fit_image(s, LOGO, (SW - 0.95) / 2, 0.9, 0.95, 1.15)
footer(s, n)
print(f'  Slide {n}: Thank you')

# ---------------------------------------------------------------- save
prs.save(OUT)
print(f'\nSaved: {OUT}')
print(f'Total slides: {len(prs.slides)}')
