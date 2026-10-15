"""Check what's on the slide layout and master."""
from pptx import Presentation

prs = Presentation(r'c:\Users\k64169500\Downloads\AIOPS-main\AIOPS-main\Presentation\ESE_Presentation (1).pptx')

# Check layout 0 (Title Slide)
layout = prs.slide_layouts[0]
print(f"Layout: {layout.name}")
print(f"Shapes on layout:")
for j, shape in enumerate(layout.shapes):
    print(f"  Shape {j}: type={shape.shape_type} name='{shape.name}'")
    print(f"    Pos: ({shape.left/914400:.2f}in, {shape.top/914400:.2f}in) Size: ({shape.width/914400:.2f}in, {shape.height/914400:.2f}in)")
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            if p.text.strip():
                print(f"    Text: '{p.text[:80]}'")
                for r in p.runs:
                    try:
                        c = r.font.color.rgb
                    except:
                        c = '(inherited)'
                    print(f"      Run: sz={r.font.size} bold={r.font.bold} name={r.font.name} color={c}")
    if shape.shape_type == 13:
        print(f"    IMAGE: {shape.image.content_type}")

print(f"\nPlaceholders on layout:")
for ph in layout.placeholders:
    print(f"  idx={ph.placeholder_format.idx} name='{ph.name}' type={ph.placeholder_format.type}")
    print(f"    Pos: ({ph.left/914400:.2f}in, {ph.top/914400:.2f}in) Size: ({ph.width/914400:.2f}in, {ph.height/914400:.2f}in)")

# Check master
master = layout.slide_master
print(f"\nMaster shapes:")
for j, shape in enumerate(master.shapes):
    print(f"  Shape {j}: type={shape.shape_type} name='{shape.name}'")
    print(f"    Pos: ({shape.left/914400:.2f}in, {shape.top/914400:.2f}in) Size: ({shape.width/914400:.2f}in, {shape.height/914400:.2f}in)")
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            if p.text.strip():
                print(f"    Text: '{p.text[:80]}'")
    if shape.shape_type == 13:
        print(f"    IMAGE: {shape.image.content_type}")

# Check what placeholders we get on a new slide
test_slide = prs.slides.add_slide(layout)
print(f"\nNew slide from layout - Shapes:")
for j, shape in enumerate(test_slide.shapes):
    print(f"  Shape {j}: type={shape.shape_type} name='{shape.name}'")
    print(f"    Pos: ({shape.left/914400:.2f}in, {shape.top/914400:.2f}in)")
print(f"\nNew slide Placeholders:")
for ph in test_slide.placeholders:
    print(f"  idx={ph.placeholder_format.idx} name='{ph.name}'")
