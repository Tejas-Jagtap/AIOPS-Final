"""
Generate the FINAL (End-Semester / Defense) presentation for the AIOps project.

Strategy (same proven approach as create_ppt.py):
  - Open the original COEP deck to preserve theme / master / fonts.
  - Delete all slides, rebuild with the FINAL-stage content.
  - Footer (COEP logo + dept text + separator + slide number) is drawn on every slide.

Difference vs the Phase-II deck (create_ppt.py):
  - Phase-III integration is now COMPLETE: FastAPI orchestrator, predictive
    autoscaler, and end-to-end validation are shown as finished work.
  - Metrics updated to the authoritative values from
    evaluation/metrics/all_metrics.json (same numbers as the journal paper).
  - New slides: Orchestrator, Predictive Autoscaler + algorithm, End-to-End
    validation, Comparison with related work, Conclusion, Publications.

Run:  python Presentation/create_final_ppt.py
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- Paths (resolve relative to this file so it works on any machine) ----
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OLD_PPT = os.path.join(BASE, 'Presentation', 'ESE_Presentation (1).pptx')
NEW_PPT = os.path.join(BASE, 'Presentation', 'AIOps_Final_Presentation.pptx')
PLOTS = os.path.join(BASE, 'evaluation', 'plots')
REPORT = os.path.join(BASE, 'Report')
LOGO_TMP = os.path.join(BASE, 'Presentation', '_logo.jpg')


def first_existing(*paths):
    """Return the first path that exists, else the last one (for warnings)."""
    for p in paths:
        if os.path.exists(p):
            return p
    return paths[-1]


ARCH_IMG = first_existing(
    os.path.join(REPORT, 'architecutre.png'),
    os.path.join(BASE, 'architecture.png'),
    os.path.join(REPORT, 'architecture.png'),
)

# ---- Step 1: Extract COEP logo from original PPT ----
src = Presentation(OLD_PPT)
for s in src.slides[0].shapes:
    if s.shape_type == 13:  # Picture
        with open(LOGO_TMP, 'wb') as f:
            f.write(s.image.blob)
        print(f"Extracted logo: {s.image.content_type}")
        break

# ---- Step 2: Open original PPT (preserves theme, master, layouts) ----
prs = Presentation(OLD_PPT)
print(f"Slide dimensions: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} inches")

# ---- Step 3: Delete all existing slides ----
sldIdLst = prs.slides._sldIdLst
for elem in list(sldIdLst):
    rId = elem.get(qn('r:id'))
    prs.part.drop_rel(rId)
    sldIdLst.remove(elem)

layout = prs.slide_layouts[0]  # "Title Slide"

# ---- Colors ----
FOOTER_CLR = RGBColor(0x20, 0x38, 0x64)
SEP_CLR    = RGBColor(0x32, 0x54, 0x90)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

# ---- Exact measurements from original PPT (EMU) ----
LOGO_L, LOGO_T  = 554040, 5662800
LOGO_W, LOGO_H  = 1283400, 1173960
FTR_L, FTR_T    = 1963440, 5648760
FTR_W, FTR_H    = 9249120, 1106542
SEP_T            = 5533200
SEP_W, SEP_H    = 12191400, 114840
SN_L, SN_T      = 8610600, 6356350
SN_W, SN_H      = 2743200, 365125
TTL_L            = 851254
TTL_W, TTL_H    = 10488891, 827882
TTL_SZ           = 508000              # 40pt
CNT_L            = 789270
CNT_T            = 1149378
CNT_W, CNT_H    = 10612859, 3835922
CNT_SZ           = 304800              # 24pt


# ===========================================================
# HELPERS
# ===========================================================

def add_footer(slide, slide_num):
    slide.shapes.add_picture(LOGO_TMP, LOGO_L, LOGO_T, LOGO_W, LOGO_H)
    tb = slide.shapes.add_textbox(FTR_L, FTR_T, FTR_W, FTR_H)
    tb.fill.background()
    tb.line.fill.background()
    tf = tb.text_frame
    bodyPr = tf._txBody.find(qn('a:bodyPr'))
    bodyPr.set('anchor', 'ctr')
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = "Dept of Computer Science and Engineering"
    r1.font.size = Pt(24); r1.font.bold = True
    r1.font.name = "Times New Roman"; r1.font.color.rgb = FOOTER_CLR
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = "COEP Technological University, Pune"
    r2.font.size = Pt(24); r2.font.bold = True
    r2.font.name = "Times New Roman"; r2.font.color.rgb = FOOTER_CLR
    sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SEP_T, SEP_W, SEP_H)
    sep.fill.solid(); sep.fill.fore_color.rgb = SEP_CLR; sep.line.color.rgb = SEP_CLR
    sn = slide.shapes.add_textbox(SN_L, SN_T, SN_W, SN_H)
    sn.text_frame.paragraphs[0].text = str(slide_num)


def set_title(slide, text, top=246660):
    ph = slide.placeholders[0]
    ph.left = TTL_L; ph.top = top; ph.width = TTL_W; ph.height = TTL_H
    p = ph.text_frame.paragraphs[0]
    p.text = text
    for run in p.runs:
        run.font.size = Emu(TTL_SZ); run.font.bold = True
    return ph


def set_content(slide, lines, font_size=CNT_SZ, align=PP_ALIGN.JUSTIFY,
                left=None, top=None, width=None, height=None):
    ph = slide.placeholders[1]
    ph.left  = left  or CNT_L
    ph.top   = top   or CNT_T
    ph.width = width or CNT_W
    ph.height = height or CNT_H
    tf = ph.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line; r.font.size = Emu(font_size)
    return ph


def remove_placeholder(slide, idx):
    if idx in slide.placeholders:
        sp = slide.placeholders[idx]._element
        sp.getparent().remove(sp)


def add_textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tb.fill.background()
    return tb.text_frame


def add_content_textbox(slide, left, top, width, height, lines,
                        font_size=CNT_SZ, align=PP_ALIGN.JUSTIFY, bold_first=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tb.fill.background()
    tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line; r.font.size = Emu(font_size)
        if bold_first and i == 0:
            r.font.bold = True
    return tf


def add_image_safe(slide, path, left, top, width, height=None):
    if os.path.exists(path):
        if height:
            slide.shapes.add_picture(path, left, top, width, height)
        else:
            slide.shapes.add_picture(path, left, top, width)
        return True
    print(f"  WARNING: not found: {path}")
    return False


def make_table(slide, rows, cols, left, top, width, height,
               headers, data, header_font=Pt(11), data_font=Pt(10)):
    ts = slide.shapes.add_table(rows, cols, left, top, width, height)
    tbl = ts.table
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        cell.text = h
        cell.fill.solid(); cell.fill.fore_color.rgb = SEP_CLR
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = header_font; run.font.bold = True
                run.font.color.rgb = WHITE
    for r, row_data in enumerate(data, 1):
        for c, val in enumerate(row_data):
            cell = tbl.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = data_font
    return tbl


# ===========================================================
# BUILD SLIDES
# ===========================================================
n = 0

# ---- SLIDE 1: Title ----
n += 1
s = prs.slides.add_slide(layout)
ph0 = s.placeholders[0]
ph0.left = TTL_L; ph0.top = 382078; ph0.width = TTL_W; ph0.height = 2078317
p = ph0.text_frame.paragraphs[0]
p.text = ("Intelligent AIOps Framework for Predictive CI/CD\n"
          "Pipeline Management and Autonomous Resource\nOptimization")
for run in p.runs:
    run.font.size = Emu(TTL_SZ); run.font.bold = True
ph1 = s.placeholders[1]
ph1.left = 1524000; ph1.top = 2888997; ph1.width = 9144000; ph1.height = 1371918
tf = ph1.text_frame
tf.paragraphs[0].text = "Tejas Annasaheb Jagtap"
tf.add_paragraph().text = "MTech Information Security"
tf.add_paragraph().text = "712442021"
gtf = add_textbox(s, 1523700, 4689517, 9144000, 570640)
r_g = gtf.paragraphs[0].add_run()
r_g.text = "Guide: Dr. V. K. Pachghare"; r_g.font.bold = True
add_footer(s, n)
print(f"  Slide {n}: Title")

# ---- SLIDE 2: Contents ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Contents", top=382079)
set_content(s, [
    "Introduction & Motivation",
    "Literature Survey & Research Gap",
    "Problem Statement & Objectives",
    "Proposed System Architecture",
    "Methodology: Data & the Three ML Models",
    "Evaluation Results (XGBoost, Autoencoder, LSTM)",
    "System Integration: FastAPI Orchestrator",
    "Predictive Autoscaler & Decision Algorithm",
    "End-to-End Validation & Demonstration",
    "Comparison with Related Work",
    "Conclusion, Future Scope & Publications",
], align=PP_ALIGN.LEFT, top=1209960)
add_footer(s, n)
print(f"  Slide {n}: Contents")

# ---- SLIDE 3: Introduction ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Introduction", top=117110)
set_content(s, [
    "Modern software delivery depends on CI/CD pipelines (Jenkins) and container orchestration (Kubernetes).",
    "These systems generate large volumes of build logs and infrastructure metrics every day.",
    "Traditional monitoring is REACTIVE - it responds only AFTER a failure or resource overload occurs:",
    "   • Build failures are discovered only after the build runs and wastes compute.",
    "   • Log issues are found by manually reading thousands of log lines.",
    "   • Kubernetes HPA scales pods only after CPU has already crossed a threshold.",
    "AIOps (AI for IT Operations) applies Machine Learning to make these operations PREDICTIVE and self-optimizing.",
])
add_footer(s, n)
print(f"  Slide {n}: Introduction")

# ---- SLIDE 4: Motivation ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Motivation")
set_content(s, [
    "Frequent build failures cause deployment delays and wasted compute resources.",
    "Manual root-cause analysis of logs is slow, tedious and error-prone.",
    "Reactive autoscaling introduces scaling lag – either over-provisioning (cost) or service degradation (SLA breach).",
    "Prior AIOps research treats these as SEPARATE problems - no unified, closed-loop system.",
    "Need: one self-learning framework that:",
    "   • Predicts build failures early from historical build patterns.",
    "   • Detects anomalies in logs automatically and in real time.",
    "   • Forecasts resource demand and scales proactively.",
    "   • Closes the loop – predictions drive automatic Kubernetes actions.",
])
add_footer(s, n)
print(f"  Slide {n}: Motivation")

# ---- SLIDE 5: Literature Survey (1) ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Literature Survey", top=206797)
remove_placeholder(s, 1)
tbl = make_table(s, 4, 4,
    Inches(0.32), Inches(1.27), Inches(12.69), Inches(4.20),
    ["Sr.", "Paper", "Work Done", "Limitation"],
    [
        ["1", "ML-Based Autoscaling for Cloud Resource Orchestration (2024)",
         "Compares ML models for predicting cluster load for autoscaling.",
         "Autoscaling only; no CI/CD or log integration."],
        ["2", "Time-series K8s Autoscaling using Prophet + LSTM (2025)",
         "Hybrid Prophet+LSTM forecast for proactive scaling.",
         "No CI/CD signals or log anomalies; scaling in isolation."],
        ["3", "Deep Learning for Multivariate Time-Series Anomaly Detection (2025)",
         "Survey of autoencoders / VAEs for anomaly detection.",
         "Survey only; no concrete DevOps pipeline application."],
    ], data_font=Pt(9))
tbl.columns[0].width = Inches(0.48); tbl.columns[1].width = Inches(2.41)
tbl.columns[2].width = Inches(4.27); tbl.columns[3].width = Inches(5.54)
add_footer(s, n)
print(f"  Slide {n}: Literature Survey (1)")

# ---- SLIDE 6: Literature Survey (2) ----
n += 1
s = prs.slides.add_slide(layout)
remove_placeholder(s, 0); remove_placeholder(s, 1)
tbl2 = make_table(s, 5, 4,
    Inches(0.12), Inches(0.11), Inches(13.12), Inches(6.10),
    ["Sr.", "Paper", "Work Done", "Limitation"],
    [
        ["4", "Towards Optimization of Anomaly Detection in DevOps (2023)",
         "Evaluates anomaly detection techniques for DevOps monitoring.",
         "No CI/CD build prediction; runtime anomalies only."],
        ["5", "AI-Enhanced Platform Engineering for CI/CD (2025)",
         "Proposes AI integration for CI/CD optimization.",
         "Conceptual; no trained models or experimental validation."],
        ["6", "Predictive Failure Analysis in CI Pipelines using Ensemble Trees (2025)",
         "Ensemble tree models for CI build failure prediction.",
         "Limited dataset; no anomaly detection or scaling."],
        ["7", "Agentic AIOps: From Reactive to Autonomous Operations (2025)",
         "Framework for closed-loop / autonomous AIOps.",
         "Conceptual closed-loop; no failure or anomaly models."],
    ], data_font=Pt(9))
tbl2.columns[0].width = Inches(0.42); tbl2.columns[1].width = Inches(2.70)
tbl2.columns[2].width = Inches(4.14); tbl2.columns[3].width = Inches(5.86)
add_footer(s, n)
print(f"  Slide {n}: Literature Survey (2)")

# ---- SLIDE 7: Research Gap ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Research Gap")
set_content(s, [
    "Existing studies address each problem in ISOLATION:",
    "   • Build-failure prediction, OR",
    "   • Log anomaly detection, OR",
    "   • Predictive autoscaling.",
    "No single framework UNIFIES all three signal types.",
    "Reactive Kubernetes HPA still dominates – predictive scaling is rare in practice.",
    "Very few works build a CLOSED LOOP: prediction → decision → automatic Kubernetes action → feedback.",
    "Gap: an integrated, closed-loop AIOps framework linking Jenkins → ML → Kubernetes.",
])
add_footer(s, n)
print(f"  Slide {n}: Research Gap")

# ---- SLIDE 8: Problem Statement ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Problem Statement", top=380148)
set_content(s, [
    '"To design and implement an intelligent, closed-loop AIOps framework that '
    'proactively predicts CI/CD pipeline failures, detects anomalies in Jenkins '
    'logs in real time, forecasts infrastructure workload, and autonomously '
    'optimizes Kubernetes resources through predictive autoscaling."',
], top=2033543, height=1421928)
add_footer(s, n)
print(f"  Slide {n}: Problem Statement")

# ---- SLIDE 9: Objectives ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Research Objectives", top=382079)
set_content(s, [
    "1. Develop a data-collection framework for CI/CD build data and cluster metrics.",
    "2. Build ML models for failure prediction, log anomaly detection and CPU forecasting.",
    "3. Implement a FastAPI-based AIOps orchestrator to serve all models via unified APIs.",
    "4. Design a predictive autoscaling engine for Kubernetes driven by ML outputs.",
    "5. Integrate the components into a closed loop and validate the system end-to-end.",
], top=1608941, height=2471446)
add_footer(s, n)
print(f"  Slide {n}: Objectives")

# ---- SLIDE 10: Architecture ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Proposed System Architecture", top=21240)
remove_placeholder(s, 1)
add_image_safe(s, ARCH_IMG, 0, 659876, 3946372, 4594637)
add_content_textbox(s, 3946372, 820000, 7266188, 4600000, [
    "Closed-loop pipeline (4 layers):",
    "1. Data Collectors",
    "     Jenkins REST API → build metadata + logs",
    "     Prometheus HTTP API → CPU / memory metrics",
    "2. ML Layer",
    "     XGBoost → build-failure prediction",
    "     TF-IDF + Autoencoder → log anomaly detection",
    "     LSTM → CPU workload forecasting",
    "3. AIOps Orchestrator (FastAPI)",
    "     Loads all models, exposes REST endpoints",
    "4. Predictive Autoscaler",
    "     Decision engine → Kubernetes Python client",
    "     Scaling actions feed back into next cycle.",
])
add_footer(s, n)
print(f"  Slide {n}: Architecture")

# ---- SLIDE 11: Methodology - Data ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Methodology: Data Pipeline & Collection", top=161736)
set_content(s, [
    "Synthetic dataset (reproducible, labelled ground truth):",
    "   • 600 Jenkins builds – 68% SUCCESS, 23% FAILURE, 9% UNSTABLE.",
    "   • 600 console log files with class-distinctive vocabulary.",
    "   • 10,080 CPU time-series points with daily / weekly patterns + noise.",
    "",
    "Live data-collection infrastructure (works with real instances):",
    "   • Jenkins REST API collector – build metadata + console logs.",
    "   • Prometheus HTTP API collector – CPU / memory metrics.",
    "   • Real-time streaming collector with configurable interval.",
    "",
    "Preprocessing: 8 build features • TF-IDF (500, n-gram 1-2) • MinMaxScaler for CPU.",
])
add_footer(s, n)
print(f"  Slide {n}: Methodology Data")

# ---- SLIDE 12: XGBoost ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Model 1: XGBoost — Failure Prediction", top=161736)
remove_placeholder(s, 1)
add_content_textbox(s, CNT_L, CNT_T, Inches(5.6), Inches(2.6), [
    "Gradient-boosted ensemble of 100 trees.",
    "Input: 8 build features (duration, warnings,",
    "  errors, failures, log_length, tests_failed/",
    "  passed/total).",
    "Output: failure probability (0-1).",
    "Class imbalance handled via scale_pos_weight.",
    "80/20 stratified split + 5-fold CV.",
    "CV F1 = 0.975 ± 0.009 (stable, no overfit).",
])
make_table(s, 6, 2, Inches(6.6), Inches(1.1), Inches(4.0), Inches(2.8),
    ["Metric", "Value"],
    [["Accuracy", "99.2%"], ["Precision", "96.4%"], ["Recall", "100%"],
     ["F1-Score", "98.2%"], ["ROC-AUC", "1.000"]])
add_image_safe(s, os.path.join(PLOTS, 'xgboost_confusion_matrix.png'), Inches(0.3), Inches(4.0), Inches(3.8))
add_image_safe(s, os.path.join(PLOTS, 'xgboost_feature_importance.png'), Inches(4.5), Inches(4.0), Inches(5.5))
add_footer(s, n)
print(f"  Slide {n}: XGBoost")

# ---- SLIDE 13: Autoencoder ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Model 2: TF-IDF + Autoencoder — Anomaly Detection", top=161736)
remove_placeholder(s, 1)
add_content_textbox(s, CNT_L, CNT_T, Inches(5.6), Inches(2.6), [
    "Two-stage, UNSUPERVISED pipeline:",
    "Stage 1: TF-IDF vectorisation (500 features,",
    "  n-gram 1-2, sublinear TF).",
    "Stage 2: Autoencoder 500→128→64→32→64→128→500.",
    "Trained on NORMAL logs only.",
    "Anomaly = high reconstruction error (MSE).",
    "Threshold = 95th percentile of normal errors.",
])
make_table(s, 6, 2, Inches(6.6), Inches(1.1), Inches(4.0), Inches(2.8),
    ["Metric", "Value"],
    [["Accuracy", "96.7%"], ["Precision", "87.1%"], ["Recall", "100%"],
     ["F1-Score", "93.1%"], ["ROC-AUC", "1.000"]])
add_image_safe(s, os.path.join(PLOTS, 'anomaly_confusion_matrix.png'), Inches(0.3), Inches(4.0), Inches(3.8))
add_image_safe(s, os.path.join(PLOTS, 'anomaly_error_distribution.png'), Inches(4.5), Inches(4.0), Inches(5.5))
add_footer(s, n)
print(f"  Slide {n}: Autoencoder")

# ---- SLIDE 14: LSTM ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Model 3: LSTM — CPU Forecasting", top=161736)
remove_placeholder(s, 1)
add_content_textbox(s, CNT_L, CNT_T, Inches(5.6), Inches(2.6), [
    "Input: 60-minute CPU history (sliding window).",
    "Output: CPU usage 10 minutes ahead.",
    "Architecture: LSTM(128) → LSTM(64) →",
    "  Dense(32, ReLU) → Dense(1).",
    "MinMaxScaler normalisation to [0,1].",
    "~10,000 training sequences, chronological split.",
    "Early stopping + LR reduction on plateau.",
])
make_table(s, 6, 2, Inches(6.6), Inches(1.1), Inches(4.0), Inches(2.6),
    ["Metric", "Value"],
    [["MSE", "0.0049"], ["MAE", "0.0433"], ["RMSE", "0.0698"],
     ["MAPE", "9.68%"], ["R²", "0.455"]])
add_image_safe(s, os.path.join(PLOTS, 'lstm_actual_vs_predicted.png'), Inches(0.3), Inches(4.0), Inches(5.5))
add_image_safe(s, os.path.join(PLOTS, 'lstm_scatter.png'), Inches(6.3), Inches(4.0), Inches(3.8))
add_footer(s, n)
print(f"  Slide {n}: LSTM")

# ---- SLIDE 15: Combined Performance ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Combined Model Performance Summary", top=161736)
remove_placeholder(s, 1)
tbl = make_table(s, 4, 4, Inches(0.5), Inches(1.3), Inches(11.5), Inches(2.0),
    ["Module", "Algorithm", "Key Metric", "ROC-AUC / R²"],
    [
        ["Failure Prediction", "XGBoost", "F1 = 0.982", "AUC = 1.000"],
        ["Anomaly Detection", "TF-IDF + Autoencoder", "F1 = 0.931", "AUC = 1.000"],
        ["CPU Forecasting", "LSTM", "RMSE = 0.070", "R² = 0.455"],
    ])
tbl.columns[0].width = Inches(2.8); tbl.columns[1].width = Inches(3.2)
tbl.columns[2].width = Inches(2.5); tbl.columns[3].width = Inches(3.0)
add_image_safe(s, os.path.join(PLOTS, 'model_comparison_summary.png'), Inches(1.0), Inches(3.5), Inches(10.0))
add_footer(s, n)
print(f"  Slide {n}: Combined Summary")

# ---- SLIDE 16: ROC Curves ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "ROC Curves — Classification Models", top=161736)
remove_placeholder(s, 1)
add_image_safe(s, os.path.join(PLOTS, 'xgboost_roc_curve.png'), Inches(0.3), Inches(1.3), Inches(5.5))
add_image_safe(s, os.path.join(PLOTS, 'anomaly_roc_curve.png'), Inches(6.3), Inches(1.3), Inches(5.5))
for x, label in [(Inches(0.3), "XGBoost: AUC = 1.000"), (Inches(6.3), "Autoencoder: AUC = 1.000")]:
    tf1 = add_textbox(s, x, Inches(5.5), Inches(5.5), Inches(0.5))
    p1 = tf1.paragraphs[0]; p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run(); r1.text = label; r1.font.size = Emu(CNT_SZ); r1.font.bold = True
add_footer(s, n)
print(f"  Slide {n}: ROC Curves")

# ---- SLIDE 17: Orchestrator ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "System Integration: FastAPI Orchestrator", top=161736)
remove_placeholder(s, 1)
add_content_textbox(s, CNT_L, Inches(1.15), Inches(6.2), Inches(4.6), [
    "Single FastAPI service that loads all three models",
    "into memory at startup and serves them via REST:",
    "",
    "   POST /predict_failure   → XGBoost",
    "   POST /detect_anomaly    → Autoencoder",
    "   POST /forecast_cpu      → LSTM",
    "   POST /autoscale         → Decision engine",
    "",
    "• Async endpoints, Pydantic request/response schemas.",
    "• Models cached in memory – millisecond inference.",
    "• Swagger UI at /docs for live testing.",
    "• Config via environment variables (no hard-coding).",
])
make_table(s, 5, 2, Inches(7.4), Inches(1.5), Inches(4.6), Inches(2.6),
    ["Endpoint", "Model"],
    [["/predict_failure", "XGBoost"],
     ["/detect_anomaly", "Autoencoder"],
     ["/forecast_cpu", "LSTM"],
     ["/autoscale", "Decision Engine"]])
add_footer(s, n)
print(f"  Slide {n}: Orchestrator")

# ---- SLIDE 18: Autoscaler + algorithm ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Predictive Autoscaler & Decision Algorithm", top=161736)
remove_placeholder(s, 1)
add_content_textbox(s, CNT_L, Inches(1.15), Inches(6.0), Inches(4.6), [
    "Priority-based decision policy (per scaling cycle):",
    "",
    "1. If anomaly detected → raise ALERT (non-blocking).",
    "2. If failure probability > 0.6 → HOLD",
    "      (do not scale a likely-failing build).",
    "3. Else if predicted CPU ≥ 0.7 → SCALE UP (+2).",
    "4. Else if predicted CPU ≤ 0.3 → SCALE DOWN (-1).",
    "5. Else → MAINTAIN.",
    "",
    "Guardrails: min 1 / max 10 replicas, 300s cooldown",
    "to prevent thrashing. Actions applied via the",
    "Kubernetes Python client → closes the loop.",
], font_size=266700)
add_content_textbox(s, Inches(7.3), Inches(1.4), Inches(4.7), Inches(4.2), [
    "Why predictive?",
    "",
    "Kubernetes HPA is REACTIVE – it scales only",
    "after CPU already crossed the threshold,",
    "causing scaling lag.",
    "",
    "Our engine acts on the FORECAST (10 min",
    "ahead), giving Kubernetes lead time to spin",
    "up pods BEFORE the spike arrives.",
], font_size=266700, bold_first=True)
add_footer(s, n)
print(f"  Slide {n}: Autoscaler")

# ---- SLIDE 19: End-to-End validation ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "End-to-End Validation & Demonstration", top=161736)
remove_placeholder(s, 1)
add_content_textbox(s, CNT_L, Inches(1.15), Inches(7.0), Inches(4.6), [
    "Automated E2E test suite: 13 / 13 tests PASS.",
    "",
    "Data layer  – dataset presence, distribution, ranges.",
    "Model layer – all 6 artefacts load correctly.",
    "API layer   – every endpoint returns correct results:",
    "   • SUCCESS build → failure prob 0.001",
    "   • FAILURE build → failure prob 0.97",
    "   • Normal log → not anomaly; Failure log → anomaly",
    "   • CPU forecast within [0,1].",
    "Full pipeline – failure + anomaly + forecast → scaling",
    "   decision in one closed-loop run.",
    "",
    "Also: Next.js demo app + Jenkinsfile (6-stage CI/CD),",
    "Docker images, and Kubernetes manifests provided.",
], font_size=266700)
make_table(s, 5, 2, Inches(8.1), Inches(1.5), Inches(3.9), Inches(2.6),
    ["Test Group", "Result"],
    [["Data layer", "3/3 PASS"],
     ["Model layer", "2/2 PASS"],
     ["API layer", "7/7 PASS"],
     ["Full pipeline", "1/1 PASS"]])
add_footer(s, n)
print(f"  Slide {n}: E2E Validation")

# ---- SLIDE 20: Comparison with related work ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Comparison with Related Work", top=161736)
remove_placeholder(s, 1)
tbl = make_table(s, 6, 5, Inches(0.4), Inches(1.4), Inches(12.5), Inches(3.2),
    ["Study", "Failure Pred.", "Log Anomaly", "Predictive Scaling", "Closed-Loop"],
    [
        ["Du et al. / Le & Zhang (log)", "No", "Yes", "No", "No"],
        ["Guruge & Priyadarshana (2025)", "No", "No", "Yes", "No"],
        ["Da Silva et al. (2025)", "No", "No", "Yes", "No"],
        ["Zota et al. (2025, agentic)", "No", "No", "No", "Yes"],
        ["Proposed Framework", "Yes", "Yes", "Yes", "Yes"],
    ], data_font=Pt(11))
tbl.columns[0].width = Inches(4.1)
add_content_textbox(s, Inches(0.5), Inches(4.9), Inches(12.0), Inches(1.0), [
    "The proposed framework is the only one integrating all four capabilities in a single closed-loop architecture.",
], font_size=266700, bold_first=True)
add_footer(s, n)
print(f"  Slide {n}: Comparison")

# ---- SLIDE 21: Conclusion ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Conclusion", top=161736)
set_content(s, [
    "Built a unified, closed-loop AIOps framework integrating three ML models with a FastAPI orchestrator and a predictive Kubernetes autoscaler.",
    "XGBoost predicts build failures with F1 = 0.982 and 100% recall (no missed failures).",
    "TF-IDF + Autoencoder detects log anomalies unsupervised with F1 = 0.931 and 100% recall.",
    "LSTM forecasts CPU with RMSE = 0.070 / MAPE = 9.68%, enabling proactive scaling.",
    "Predictive autoscaling acts on forecasts, surpassing the reactive Kubernetes HPA.",
    "Validated end-to-end (13/13 tests) - a concrete implementation of the closed-loop AIOps advocated in recent literature.",
])
add_footer(s, n)
print(f"  Slide {n}: Conclusion")

# ---- SLIDE 22: Future Scope ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Future Scope", top=161736)
set_content(s, [
    "Validate on real-world CI/CD datasets (e.g., TravisTorrent, BuildFast) and live Jenkins/Prometheus data.",
    "Add automatic retraining to handle concept drift as pipelines evolve.",
    "Replace TF-IDF with transformer-based log embeddings (e.g., LogBERT) for richer context.",
    "Benchmark the predictive autoscaler against Kubernetes HPA under production-scale load.",
    "Build a monitoring dashboard for predictions, alerts and scaling actions.",
    "Extend to multi-service dependency and memory / network forecasting.",
], top=1024061, height=4154984)
add_footer(s, n)
print(f"  Slide {n}: Future Scope")

# ---- SLIDE 23: Publications ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "Publications", top=161736)
set_content(s, [
    "The work has been documented as a research manuscript (Springer Nature journal format), with a conference-paper version also prepared.",
    "",
    "Title: \"Intelligent AIOps Framework for Predictive CI/CD Pipeline Management and Autonomous Resource Optimization.\"",
    "",
    "Authors: Tejas A. Jagtap, Charulata Patil, V. K. Pachghare (Dept. of Computer Science & Engineering, COEP Technological University, Pune).",
    "",
    "Contribution: first unified closed-loop AIOps framework combining failure prediction, log anomaly detection, CPU forecasting and predictive autoscaling.",
], top=1024061, height=4154984)
add_footer(s, n)
print(f"  Slide {n}: Publications")

# ---- SLIDE 24: References ----
n += 1
s = prs.slides.add_slide(layout)
set_title(s, "References", top=382079)
set_content(s, [
    '[1] I. Pintye et al., "Enhancing ML-Based Autoscaling for Cloud Resource Orchestration," J. Grid Computing, 2024.',
    '[2] P. B. Guruge and Y. H. P. P. Priyadarshana, "Time-Series Forecasting-Based Kubernetes Autoscaling Using Prophet and LSTM," Frontiers in CS, 2025.',
    '[3] M. A. Yahya et al., "Deep Learning for Multivariate Time-Series Anomaly Detection: A Survey," AI Review, 2025.',
    '[4] A. Hrusto, "Towards Optimization of Anomaly Detection in DevOps," IST, 2023.',
    '[5] P. Mathew and A. Baswareddy, "Predictive Failure Analysis in CI Pipelines Using Ensemble Trees," IJACS, 2025.',
    '[6] M. Du et al., "DeepLog: Anomaly Detection from System Logs," CCS, 2017.',
    '[7] Y. Zota et al., "A Practical Framework for Agentic AIOps," 2025.',
    '[8] Kubernetes Documentation, "Horizontal Pod Autoscaler," 2024.',
], font_size=254000, top=1438739, height=3652282)
add_footer(s, n)
print(f"  Slide {n}: References")

# ---- SLIDE 25: Thank You ----
n += 1
s = prs.slides.add_slide(layout)
remove_placeholder(s, 1)
ph = s.placeholders[0]
ph.left = TTL_L; ph.top = 2186339; ph.width = TTL_W; ph.height = TTL_H
p = ph.text_frame.paragraphs[0]
p.text = "Thank You !!!"
for run in p.runs:
    run.font.size = Pt(72); run.font.bold = True
add_footer(s, n)
print(f"  Slide {n}: Thank You")

# ---- SAVE ----
prs.save(NEW_PPT)
print(f"\nPresentation saved to: {NEW_PPT}")
print(f"Total slides: {len(prs.slides)}")
try:
    os.remove(LOGO_TMP)
except OSError:
    pass
