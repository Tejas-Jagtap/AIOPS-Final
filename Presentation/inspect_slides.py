"""Extract first 6 slides in detail."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu

prs = Presentation(r'c:\Users\k64169500\Downloads\AIOPS-main\AIOPS-main\Presentation\ESE_Presentation (1).pptx')

print(f"Slide width: {prs.slide_width} EMU = {prs.slide_width/914400:.2f} inches")
print(f"Slide height: {prs.slide_height} EMU = {prs.slide_height/914400:.2f} inches")
print(f"Total slides: {len(prs.slides)}")
print(f"Available layouts: {[(i, l.name) for i, l in enumerate(prs.slide_layouts)]}")
print()

for i, slide in enumerate(prs.slides):
    if i >= 6:
        break
    print(f"{'='*60}")
    print(f"SLIDE {i+1} - Layout: {slide.slide_layout.name}")
    print(f"{'='*60}")
    
    for j, shape in enumerate(slide.shapes):
        print(f"\n  Shape {j}: type={shape.shape_type}, name='{shape.name}'")
        print(f"    Pos: left={shape.left}({shape.left/914400:.2f}in) top={shape.top}({shape.top/914400:.2f}in)")
        print(f"    Size: w={shape.width}({shape.width/914400:.2f}in) h={shape.height}({shape.height/914400:.2f}in)")
        
        if hasattr(shape, 'fill'):
            try:
                f = shape.fill
                if f.type is not None:
                    print(f"    Fill: type={f.type}", end='')
                    if f.type == 1:  # Solid
                        try:
                            print(f" color={f.fore_color.rgb}")
                        except:
                            print(" color=(theme)")
                    else:
                        print()
            except:
                pass
        
        if shape.shape_type == 13:
            print(f"    IMAGE: {shape.image.content_type}")
        
        if shape.has_text_frame:
            tf = shape.text_frame
            for k, para in enumerate(tf.paragraphs):
                if para.text.strip():
                    align = para.alignment
                    print(f"    P{k}: '{para.text[:100]}'  align={align}")
                    for r, run in enumerate(para.runs):
                        if run.text.strip():
                            sz = run.font.size
                            b = run.font.bold
                            nm = run.font.name
                            try:
                                c = run.font.color.rgb
                            except:
                                c = '(inherited)'
                            print(f"      R{r}: sz={sz} bold={b} name={nm} color={c}")
        
        if shape.has_table:
            t = shape.table
            print(f"    TABLE: {len(t.rows)}x{len(t.columns)}")
            print(f"    Col widths: {[f'{c.width/914400:.2f}' for c in t.columns]}")
            # first row
            for ci in range(len(t.columns)):
                cell = t.cell(0, ci)
                print(f"      Header[{ci}]: '{cell.text}'", end='')
                try:
                    print(f" fill={cell.fill.fore_color.rgb}", end='')
                except:
                    pass
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        try:
                            print(f" sz={run.font.size} color={run.font.color.rgb}", end='')
                        except:
                            pass
                print()
    print()
