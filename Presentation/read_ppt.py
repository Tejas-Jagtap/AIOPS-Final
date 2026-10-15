from pptx import Presentation
import sys

prs = Presentation(r'c:\Users\k64169500\Downloads\AIOPS-main\AIOPS-main\Presentation\ESE_Presentation (1).pptx')
for i, slide in enumerate(prs.slides):
    print(f'=== SLIDE {i+1} ===')
    print(f'Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    print(f'  {text}')
        if shape.has_table:
            table = shape.table
            for row in table.rows:
                row_text = ' | '.join([cell.text.strip() for cell in row.cells])
                print(f'  | {row_text} |')
        if shape.shape_type == 13:
            print(f'  [IMAGE: {shape.image.content_type}]')
    print()
